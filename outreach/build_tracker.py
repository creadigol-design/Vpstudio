import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Outreach Tracker"

# ---- palette ----
navy = "111118"
green = "9FCC3B"
lightgreen = "EAF4CE"
grey = "F2F2F2"
white = "FFFFFF"

thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = [
    "#", "Brand", "Category", "Location", "Region", "The gap (evidence)",
    "Funding / growth signal", "Website", "Instagram", "Followers",
    "Contact name", "Contact route", "Verified email", "Email subject",
    "Draft ready?", "Reply status", "Notes",
]

rows = [
    [1, "Snowdonia Cheese Company", "Food — cheese", "Rhyl, Denbighshire", "North Wales ⭐",
     "~34k IG, product-shot heavy; iconic Black Bomber barely used in video",
     "Queen's Award (Int'l Trade); exports 25+ countries; ~113 awards", "snowdoniacheese.co.uk",
     "@snowdoniacheese", "~34k", "John & Richard Newton-Jones",
     "LinkedIn / site contact form", "", "Black Bomber deserves more than a product shot",
     "Yes – add recipient", "Not sent", ""],
    [2, "Wild Horse Brewing Co", "Drink — craft beer", "Llandudno, Conwy", "North Wales ⭐",
     "~7.5k IG / ~1,287 posts; lots of effort, little reach; design-led cans unused",
     "DBW/Welsh Gov funding; capacity doubled 175k→350k L; new taproom", "wildhorsebrewing.co.uk",
     "@wildhorsebeer", "~7.5k", "Dave & Emma Faragher",
     "LinkedIn / site contact form", "", "Those cans are wasted as thumbnails",
     "Yes – add recipient", "Not sent", "Warmest travel pitch – on the doorstep"],
    [3, "Different Dog", "Pet — fresh dog food", "Shrewsbury, Shropshire", "Welsh border",
     "~43k IG, active but product/testimonial-led; light for the funding",
     "~£10m raised (2025), earmarked for marketing; 18,000+ dogs/mo; B-Corp", "differentdog.com",
     "@differentdog", "~43k", "Charlie & Alex Thurstan",
     "LinkedIn / site contact form", "", "£10m in, and the dogs aren't on camera enough",
     "Yes – add recipient", "Not sent", "~1 hr from DocShed"],
    [4, "Hilltop Honey", "Food — honey", "Newtown, Powys", "Mid Wales",
     "~39k IG / ~1,524 posts; jar photography heavy; lags commercial scale",
     "£10m Santander (2024); sales £20m→£33m, target £50m; B-Corp; major multiples", "hilltophoney.co.uk",
     "@hilltop_honey", "~39k", "Scott Davies (CEO)",
     "LinkedIn / site contact form", "", "£33m in sales, 39k followers — the maths is off",
     "Yes – add recipient", "Not sent", ""],
    [5, "Hip Pop", "Drink — gut soda", "Altrincham, Gtr Manchester", "North West",
     "~29k IG; invests in content but underweight vs US peers (millions)",
     "~£7m raised (LadBible founders, ex-Olipop CFO); Waitrose & Booths", "drinkhippop.com",
     "@drink.hip.pop", "~29k", "Emma Thackray & Kenny Goodman",
     "hello@drinkhippop.com", "hello@drinkhippop.com", "29k is underweight for where Hip Pop's headed",
     "YES – ready to send", "Not sent", "Verified email – send-ready"],
    [6, "The Forest Distillery (Forest Gin)", "Drink — spirits", "Macclesfield, Cheshire", "North West",
     "~16k IG; premium, cinematic product with almost no organic video",
     "Double double-gold SFWSC; 'World's Best 50 Gins'; Harvey Nichols", "theforestdistillery.com",
     "@forest_distillery", "~16k", "Karl & Lindsay Bond",
     "LinkedIn / site contact form", "", "The most beautiful bottle in Britain, hiding at 16k",
     "Yes – add recipient", "Not sent", "Most visually stunning product on the list"],
    [7, "Adamo Foods", "Foodtech — mycelium steak", "London", "National",
     "~995 IG followers / ~19 posts; near-absent owned audience",
     "~£2m seed + €10m EU Horizon grant; consumer launch ~2027 (build-hype angle)", "adamofoods.com",
     "@adamo.foods", "~995", "Pierre Dupuis (CEO)",
     "LinkedIn; UNVERIFIED contact@/marketing@adamofoods.com", "", "€10m raised, 995 followers — that gap is the opportunity",
     "Yes – add recipient", "Not sent", "Highest-contrast story; verify email first"],
    [8, "PURIFIED", "Footwear", "London", "National",
     "~3.4k IG; world-first plastic-free shoe + royal exposure, static feed",
     "Prince William wore them at Earthshot 2024; raising investment", "purified.eco",
     "@purifiedfootwear", "~3.4k", "Will Verona (founder)",
     "enquiries@purified.eco (press: charlie@purified.eco)", "enquiries@purified.eco", "Prince William wore your shoes — where's the video?",
     "YES – ready to send", "Not sent", "Verified email – send-ready"],
    [9, "ReBorn Homewares", "Homeware", "Wiltshire", "National",
     "~6.4k IG; styled static photography, no video engine",
     "Holly Branson + Conduit EIS backing; ~$550k angel (2025); John Lewis", "reborn.homes",
     "@reborn_homes", "~6.4k", "Brian Walmsley (CEO)",
     "LinkedIn (in/brianwalmsley) / site form", "", "Waste into homeware — that's a film, not a photo",
     "Yes – add recipient", "Not sent", "Waste-to-product = ready-made video"],
    [10, "ACTIPH Water", "Drink — water", "London", "National",
     "IG ~26k but TikTok ~1,750 — dormant on the key beverage channel",
     "~£10m+ raised; Tesco/Sainsbury's/Ocado/H&B/BP; exports ~20 countries", "actiphwater.com",
     "@actiphwater", "TikTok ~1,750", "Barnaby Hughes (CMO) / Jamie Douglas-Hamilton",
     "Barnaby: likely firstname@actiphwater.com (verify)", "", "Your TikTok's asleep — and it's costing you",
     "Yes – add recipient", "Not sent", ""],
]

# ---- title ----
ws.merge_cells("A1:Q1")
c = ws["A1"]
c.value = "vedrí — Outreach Lead Tracker  ·  10 brands  ·  prepared 1 Aug 2026"
c.font = Font(name="Arial", size=14, bold=True, color=white)
c.fill = PatternFill("solid", fgColor=navy)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[1].height = 30

# ---- legend ----
ws.merge_cells("A2:Q2")
c = ws["A2"]
c.value = ("Legend:  Green rows = verified email, send-ready.  "
           "All other drafts are saved in Gmail with your own address as a placeholder — replace the To before sending.  "
           "Follower counts are search-surfaced (Aug 2026) — confirm on the live feed before quoting in line one.  "
           "Fill in the 'Reply status' and 'Notes' columns as you work through.")
c.font = Font(name="Arial", size=9, italic=True, color="444444")
c.fill = PatternFill("solid", fgColor=lightgreen)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
ws.row_dimensions[2].height = 40

# ---- header row (row 4) ----
hrow = 4
for j, h in enumerate(headers, start=1):
    cell = ws.cell(row=hrow, column=j, value=h)
    cell.font = Font(name="Arial", size=10, bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[hrow].height = 28

# ---- data rows ----
verified_ranks = {5, 8}
for i, row in enumerate(rows):
    r = hrow + 1 + i
    fill_color = lightgreen if row[0] in verified_ranks else (white if i % 2 == 0 else grey)
    for j, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=j, value=val)
        cell.font = Font(name="Arial", size=9, color=navy)
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = border
    ws.row_dimensions[r].height = 58

# ---- column widths ----
widths = [4, 22, 18, 20, 13, 34, 34, 20, 17, 12, 22, 26, 30, 32, 16, 14, 24]
for j, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

ws.freeze_panes = "C5"

out = "/home/user/Vpstudio/outreach/vedri-outreach-tracker.xlsx"
wb.save(out)
print("saved", out)
